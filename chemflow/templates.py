"""
Generator template Excel input chemflow (ligan & reseptor).

Dipanggil dari ``chemflow init`` sehingga contoh input selalu konsisten
dengan kolom yang benar-benar dibaca ``io.excel_ligands``/``io.excel_receptors``
saat itu, tanpa perlu menyimpan file contoh statis di repo yang bisa
kadaluarsa begitu skema kolom berubah.
"""

from __future__ import annotations

from pathlib import Path
from typing import List, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_FONT_NAME = "Arial"

_F_HEADER_WAJIB = Font(name=_FONT_NAME, size=10, bold=True, color="FFFFFF")
_F_HEADER_OPSIONAL = Font(name=_FONT_NAME, size=10, bold=True, color="000000")
_F_BODY = Font(name=_FONT_NAME, size=10)
_F_NOTE = Font(name=_FONT_NAME, size=9, italic=True, color="595959")
_F_SECTION = Font(name=_FONT_NAME, size=11, bold=True, color="1F4E78")
_F_TITLE = Font(name=_FONT_NAME, size=13, bold=True, color="FFFFFF")
_F_SUBTITLE = Font(name=_FONT_NAME, size=10, italic=True, color="595959")

_FILL_WAJIB = PatternFill("solid", fgColor="C00000")
_FILL_OPSIONAL = PatternFill("solid", fgColor="D9E1F2")
_FILL_EXAMPLE = PatternFill("solid", fgColor="F2F2F2")
_FILL_TITLE = PatternFill("solid", fgColor="1F4E78")

_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _header_cell(ws: Worksheet, row: int, col: int, text: str, wajib: bool) -> None:
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = _F_HEADER_WAJIB if wajib else _F_HEADER_OPSIONAL
    cell.fill = _FILL_WAJIB if wajib else _FILL_OPSIONAL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = _BORDER


def _body_cell(ws: Worksheet, row: int, col: int, value) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = _F_BODY
    cell.border = _BORDER
    cell.fill = _FILL_EXAMPLE
    cell.alignment = Alignment(horizontal="left", vertical="center")


def _set_widths(ws: Worksheet, widths: Sequence[float]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _add_instructions_sheet(
    wb: Workbook,
    subtitle: str,
    keyword_rows: List[Tuple[str, str, str, str]],
    general_notes: List[str],
) -> None:
    ws = wb.create_sheet("Petunjuk")
    ws.sheet_view.showGridLines = False

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
    title_cell = ws.cell(row=1, column=1, value=f"chemflow: {subtitle}")
    title_cell.font = _F_TITLE
    title_cell.fill = _FILL_TITLE
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 26

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=5)
    ws.cell(row=2, column=1,
            value="Sheet data ada di tab pertama. Sheet ini hanya referensi kolom.").font = _F_SUBTITLE

    r = 4
    ws.cell(row=r, column=1, value="Legenda warna header").font = _F_SECTION
    r += 1
    cell = ws.cell(row=r, column=1, value="WAJIB DIISI")
    cell.font = _F_HEADER_WAJIB
    cell.fill = _FILL_WAJIB
    cell.alignment = Alignment(horizontal="center")
    cell.border = _BORDER
    ws.cell(row=r, column=2, value="Kolom harus ada dan berisi nilai di setiap baris.").font = _F_BODY
    r += 1
    cell = ws.cell(row=r, column=1, value="OPSIONAL")
    cell.font = _F_HEADER_OPSIONAL
    cell.fill = _FILL_OPSIONAL
    cell.alignment = Alignment(horizontal="center")
    cell.border = _BORDER
    ws.cell(row=r, column=2, value="Boleh dikosongkan atau kolom boleh dihapus. chemflow pakai nilai default.").font = _F_BODY

    r += 3
    ws.cell(row=r, column=1, value="Nama kolom yang dikenali otomatis (case-insensitive, cukup mengandung kata kunci)").font = _F_SECTION
    r += 1
    for i, h in enumerate(("Field", "Wajib?", "Kata kunci yang dikenali", "Keterangan"), start=1):
        _header_cell(ws, r, i, h, wajib=False)
    for row_data in keyword_rows:
        r += 1
        for i, val in enumerate(row_data, start=1):
            cell = ws.cell(row=r, column=i, value=val)
            cell.font = Font(name=_FONT_NAME, size=10, bold=True, color="C00000") if (i == 2 and val == "Wajib") else _F_BODY
            cell.border = _BORDER
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    r += 3
    ws.cell(row=r, column=1, value="Catatan").font = _F_SECTION
    r += 1
    for note in general_notes:
        ws.cell(row=r, column=1, value=f"- {note}").font = _F_BODY
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        ws.row_dimensions[r].height = 28
        r += 1

    _set_widths(ws, (18, 10, 30, 50))


def write_ligand_tidy_template(path: Path) -> Path:
    """Tulis template ligan format tidy (1 baris = 1 senyawa) siap dipakai
    langsung sebagai ``--ligands``."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Ligan"
    ws.sheet_view.showGridLines = False

    _header_cell(ws, 1, 1, "name", wajib=True)
    _header_cell(ws, 1, 2, "smiles", wajib=False)
    _header_cell(ws, 1, 3, "group", wajib=False)

    examples = [
        ("Quercetin", "Oc1cc(O)c2c(=O)c(O)c(-c3ccc(O)c(O)c3)oc2c1", "Tanaman_X"),
        ("Curcumin", "COc1cc(/C=C/C(=O)CC(=O)/C=C/c2ccc(O)c(OC)c2)ccc1O", "Tanaman_Y"),
        ("Aspirin", "", "Obat_Pembanding"),
        ("Kaempferol", "O=c1c(O)c(-c2ccc(O)cc2)oc2cc(O)cc(O)c12", ""),
    ]
    for i, row_vals in enumerate(examples):
        for j, val in enumerate(row_vals, start=1):
            _body_cell(ws, 2 + i, j, val)

    _set_widths(ws, (22, 55, 20))

    _add_instructions_sheet(
        wb, "Ligan (format Tidy)",
        keyword_rows=[
            ("name", "Wajib", "compound name, name", "Nama bebas, dipakai sebagai nama file & label hasil."),
            ("smiles", "Opsional", "smiles", "Kosongkan agar SMILES dicari otomatis via PubChem. Kehadiran kolom ini (walau sebagian sel kosong) membuat file dibaca sebagai format Tidy, bukan Wide."),
            ("group", "Opsional", "group, sumber, source", "Label kelompok sumber senyawa, dipakai untuk PCA kemometrik (perlu minimal 2 kelompok berbeda di seluruh file)."),
        ],
        general_notes=[
            "Baris \"Aspirin\" sengaja mengosongkan SMILES sebagai contoh. chemflow akan mencarikannya otomatis via PubChem, atau menawarkan input manual jika PubChem gagal.",
            "Ada juga format Wide (tanpa kolom smiles, tiap kolom = 1 kelompok). Buat dengan: chemflow init --format wide.",
            "Jika memakai file hasil ADMETLab3 (--admet-file), baris file itu harus berurutan sama dengan baris sheet ini dari atas ke bawah.",
        ],
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path


def write_ligand_wide_template(path: Path) -> Path:
    """Tulis template ligan format wide (tiap kolom = 1 kelompok) siap
    dipakai langsung sebagai ``--ligands``."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Ligan"
    ws.sheet_view.showGridLines = False

    groups = ["Tanaman_X", "Tanaman_Y", "Tanaman_Z"]
    for i, g in enumerate(groups, start=1):
        _header_cell(ws, 1, i, g, wajib=False)

    rows = [
        ["Quercetin", "Curcumin", "Gingerol"],
        ["Kaempferol", "Demethoxycurcumin", "Shogaol"],
        ["Rutin", "", "Zingerone"],
    ]
    for i, row_vals in enumerate(rows):
        for j, val in enumerate(row_vals, start=1):
            if val:
                _body_cell(ws, 2 + i, j, val)

    _set_widths(ws, (22, 22, 22))

    _add_instructions_sheet(
        wb, "Ligan (format Wide)",
        keyword_rows=[
            ("(header kolom)", "-", "bebas, header = nama kelompok", "Tiap kolom mewakili 1 kelompok sumber senyawa. Isi sel = nama senyawa dalam kelompok itu."),
        ],
        general_notes=[
            "Format ini dipakai hanya jika tidak ADA kolom bernama \"smiles\" di seluruh sheet.",
            "Panjang tiap kolom boleh beda (sel kosong diabaikan).",
            "Semua nama di sini dicari SMILES-nya otomatis via PubChem (format ini tidak punya kolom SMILES sama sekali).",
            "Jika memakai file hasil ADMETLab3 (--admet-file), urutan ligan adalah kolom dari kiri ke kanan, tiap kolom dari atas ke bawah.",
            "Ada juga format Tidy (1 baris = 1 senyawa + kolom smiles opsional). Buat dengan: chemflow init --format tidy.",
        ],
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path


def write_receptor_template(path: Path) -> Path:
    """Tulis template reseptor siap dipakai langsung sebagai ``--receptors``."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Reseptor"
    ws.sheet_view.showGridLines = False

    headers = [
        ("pdb_code", True), ("center_x", False), ("center_y", False), ("center_z", False),
        ("size_x", False), ("size_y", False), ("size_z", False),
    ]
    for i, (h, wajib) in enumerate(headers, start=1):
        _header_cell(ws, 1, i, h, wajib)

    examples = [
        ("6LU7", "", "", "", "", "", ""),
        ("3PTB", -1.76, 14.46, 16.92, 20, 20, 20),
        ("3PTB", 10.2, 5.1, -3.4, 18, 18, 18),
        ("1AKI", "", "", "", "", "", ""),
    ]
    for i, row_vals in enumerate(examples):
        for j, val in enumerate(row_vals, start=1):
            _body_cell(ws, 2 + i, j, val)

    _set_widths(ws, (12, 11, 11, 11, 10, 10, 10))

    _add_instructions_sheet(
        wb, "Reseptor",
        keyword_rows=[
            ("pdb_code", "Wajib", "pdb code, pdb", "Kode 4 karakter RCSB PDB, mis. 6LU7. File .pdb diunduh otomatis."),
            ("center_x", "Opsional", "center_x, cx", "Koordinat Angstrom pusat kotak pencarian docking."),
            ("center_y", "Opsional", "center_y, cy", "Idem, sumbu Y."),
            ("center_z", "Opsional", "center_z, cz", "Idem, sumbu Z."),
            ("size_x", "Opsional", "size_x, sx", "Lebar kotak pencarian (Angstrom), default 20 jika kosong."),
            ("size_y", "Opsional", "size_y, sy", "Idem, sumbu Y."),
            ("size_z", "Opsional", "size_z, sz", "Idem, sumbu Z."),
            ("(alternatif ukuran)", "Opsional", "gridbox, box_size", "1 kolom untuk size_x=size_y=size_z sekaligus, alternatif dari 3 kolom size_x/y/z terpisah."),
        ],
        general_notes=[
            "Baris 6LU7: gridbox kosong. chemflow mendeteksi ligan native dari struktur PDB dan meminta Anda pilih interaktif saat pipeline berjalan.",
            "Baris 3PTB muncul 2x dengan gridbox berbeda: 2 titik docking (multi-situs) pada reseptor yang sama.",
            "Kode PDB berbeda (6LU7, 3PTB, 1AKI): otomatis docking multi-reseptor, semua ligan didock ke semua reseptor.",
            "Belum tahu koordinat dan ukuran gridbox? Jalankan: chemflow receptor-config. Perintah itu menanyakan kode PDB dan ligan native, lalu menulis Excel reseptor lengkap.",
        ],
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path


def write_all_templates(output_dir: Path) -> List[Path]:
    """Tulis ketiga template (ligan tidy, ligan wide, reseptor) ke satu direktori.

    Args:
        output_dir: direktori tujuan, dibuat otomatis jika belum ada.

    Returns:
        Daftar path file yang ditulis.
    """
    output_dir.mkdir(parents=True, exist_ok=True)
    return [
        write_ligand_tidy_template(output_dir / "ligan_contoh_tidy.xlsx"),
        write_ligand_wide_template(output_dir / "ligan_contoh_wide.xlsx"),
        write_receptor_template(output_dir / "reseptor_contoh.xlsx"),
    ]
