"""
Pembaca file Excel interaksi ligan-reseptor dari BIOVIA Discovery Studio.

Sumbernya tabel Non-bond BIOVIA, baik hasil ``chemflow.interaction`` (otomatis)
maupun ekspor manual. Kolom yang dipakai: Category, Types (atau Type/Subtype),
From, From Chemistry, To, To Chemistry. Bila baris pertama berisi judul kolom,
kolom dicari lewat namanya, jadi urutan dan kolom tambahan (Distance, Angle,
ID, ...) tidak berpengaruh. Tanpa judul, dipakai urutan posisi bawaan
Discovery Studio: Name, Visible/Rendered, Color, Parent/Style, Distance/ID,
Category, Types, From, From Chemistry, To, To Chemistry (indeks 0 sampai 10).

Setiap baris adalah satu interaksi antara dua sisi. Sisi ditulis sebagai
"Chain:RESNAME+RESNUM:ATOMNAME" (mis. "A:SER195:OG"). Interaksi pi
(Pi-Sigma, Pi-Pi) menuliskan cincin aromatik tanpa nama atom ("A:PHE330"),
sehingga bagian atom bersifat opsional.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import List

_MIN_COLUMNS = 11  # tanpa judul kolom: Name..To Chemistry (indeks 0-10)
_POSITIONAL = {"category": 5, "type_subtype": 6, "from": 7, "from_chemistry": 8, "to": 9, "to_chemistry": 10}
_HEADER_NAMES = {
    "category": ("category",),
    "type_subtype": ("types", "type", "type/subtype"),
    "from": ("from",),
    "from_chemistry": ("from chemistry",),
    "to": ("to",),
    "to_chemistry": ("to chemistry",),
}

_ATOM_SPEC_RE = re.compile(
    r"^\s*(?P<chain>[A-Za-z0-9]{1,2}):(?P<resname>[A-Za-z0-9]{1,4}?)(?P<resnum>-?\d+)(?::(?P<atom>.+?))?\s*$"
)


@dataclass
class AtomSpec:
    """Satu sisi interaksi (hasil parse "Chain:RESNAME+RESNUM:ATOMNAME")."""
    chain: str
    resname: str
    resnum: int
    atom_name: str

    @property
    def residue_key(self) -> tuple:
        return (self.resnum, self.resname)

    @property
    def residue_label(self) -> str:
        return f"{self.resnum}-{self.resname.title()}"


@dataclass
class BiovaInteraction:
    """Satu baris interaksi dari export BIOVIA (kolom sampai To Chemistry)."""
    category: str
    type_subtype: str
    from_spec: AtomSpec
    from_chemistry: str
    to_spec: AtomSpec
    to_chemistry: str


def parse_atom_spec(raw: object) -> AtomSpec:
    """Parse spek atom BIOVIA "Chain:RESNAME+RESNUM:ATOMNAME".

    Raises:
        ValueError: format tidak cocok pola yang diharapkan.
    """
    text = str(raw or "").strip()
    match = _ATOM_SPEC_RE.match(text)
    if not match:
        raise ValueError(f"Format spek atom BIOVIA tidak dikenali: {raw!r}")
    return AtomSpec(
        chain=match.group("chain"), resname=match.group("resname").upper(),
        resnum=int(match.group("resnum")), atom_name=match.group("atom") or "",
    )


def read_biovia_interactions(path: "str | Path") -> List[BiovaInteraction]:
    """Baca file Excel interaksi BIOVIA Discovery Studio.

    Bila baris pertama berisi judul kolom (From, To, Category, Types), kolom
    dicari lewat namanya dan baris itu dilewati; file yang hanya berisi judul
    berarti kompleks tanpa interaksi dan menghasilkan daftar kosong. Tanpa
    judul dipakai urutan posisi bawaan (lihat docstring modul). Baris pertama
    tanpa judul yang dikenali dianggap header hanya bila From/To-nya bukan
    spek atom; jika berupa spek atom, seluruh baris dianggap data.

    Args:
        path: path file .xlsx interaksi BIOVIA.

    Returns:
        Daftar ``BiovaInteraction``, satu per baris data.

    Raises:
        FileNotFoundError: file tidak ditemukan.
        ValueError: file kosong, kolom wajib tidak lengkap, atau spek atom
            From/To tidak bisa diparse (pesan menyebut nama file dan nomor baris).
    """
    from openpyxl import load_workbook

    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"File interaksi BIOVIA tidak ditemukan: {path}")

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.worksheets[0]

    rows = [list(row) for row in ws.iter_rows(values_only=True) if any(c is not None for c in row)]
    wb.close()

    if not rows:
        raise ValueError(f"File interaksi BIOVIA kosong: {path}")

    columns = _columns_from_header(rows[0])
    start = 0
    if columns is not None:
        start = 1
    elif len(rows) > 1 and not _looks_like_data_row(rows[0]):
        start = 1  # baris pertama judul yang tak dikenali; hanya dilewati bila ada baris data setelahnya
    if columns is None:
        columns = _POSITIONAL
    needed = max(columns["from"], columns["to"], columns["category"], columns["type_subtype"]) + 1
    if columns is _POSITIONAL:
        needed = _MIN_COLUMNS

    def cell(row: list, key: str) -> str:
        index = columns.get(key)
        return str(row[index] or "") if index is not None and index < len(row) else ""

    interactions: List[BiovaInteraction] = []
    for row_idx, row in enumerate(rows[start:], start=start + 1):
        if len(row) < needed:
            raise ValueError(
                f"Baris {row_idx} di '{path.name}' hanya punya {len(row)} kolom, minimal {needed} "
                f"diperlukan. Pastikan file adalah export BIOVIA Discovery Studio yang valid."
            )
        try:
            from_spec = parse_atom_spec(cell(row, "from"))
            to_spec = parse_atom_spec(cell(row, "to"))
        except ValueError as exc:
            raise ValueError(f"Baris {row_idx} di '{path.name}': {exc}") from exc

        interactions.append(BiovaInteraction(
            category=cell(row, "category"), type_subtype=cell(row, "type_subtype"),
            from_spec=from_spec, from_chemistry=cell(row, "from_chemistry"),
            to_spec=to_spec, to_chemistry=cell(row, "to_chemistry"),
        ))

    return interactions


def _columns_from_header(row: list) -> "dict | None":
    """Peta kolom dari baris judul, atau None bila baris itu bukan judul Non-bond BIOVIA."""
    names = [str(cell or "").strip().lower() for cell in row]
    columns = {}
    for key, aliases in _HEADER_NAMES.items():
        for alias in aliases:
            if alias in names:
                columns[key] = names.index(alias)
                break
    return columns if {"category", "type_subtype", "from", "to"} <= columns.keys() else None


def _looks_like_data_row(row) -> bool:
    if len(row) < _MIN_COLUMNS:
        return False
    try:
        parse_atom_spec(row[7])
        parse_atom_spec(row[9])
        return True
    except ValueError:
        return False


@dataclass
class ProteinLigandContact:
    """Satu kontak protein-ligan (sisi protein + tipe interaksi)."""
    resnum: int
    resname: str
    interaction_type: str

    @property
    def residue_key(self) -> tuple:
        return (self.resnum, self.resname)

    @property
    def residue_label(self) -> str:
        return f"{self.resnum}-{self.resname.title()}"


def filter_protein_ligand(interactions: List[BiovaInteraction], ligand_chain: str) -> List[ProteinLigandContact]:
    """Saring interaksi murni protein-ligan (buang intra-ligand & intra-protein).

    Args:
        interactions: hasil ``read_biovia_interactions()``.
        ligand_chain: chain ID ligan (dari metadata sidecar merge, mis. "X").

    Returns:
        Daftar ``ProteinLigandContact``, sisi protein saja (chain != ligand_chain).
    """
    contacts: List[ProteinLigandContact] = []
    for it in interactions:
        from_is_ligand = it.from_spec.chain == ligand_chain
        to_is_ligand = it.to_spec.chain == ligand_chain
        if from_is_ligand == to_is_ligand:
            continue  # keduanya ligan (intra-ligand) atau keduanya protein (intra-protein)

        protein_spec = it.to_spec if from_is_ligand else it.from_spec
        contacts.append(ProteinLigandContact(
            resnum=protein_spec.resnum, resname=protein_spec.resname, interaction_type=it.type_subtype,
        ))
    return contacts
