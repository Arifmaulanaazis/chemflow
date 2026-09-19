"""Utilitas Excel bersama: gaya header, freeze panes, filter, dan autofit kolom."""

from __future__ import annotations

from pathlib import Path
from typing import Dict

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT_NAME, size=10)


def style_header(ws: Worksheet) -> None:
    """Header tebal berlatar biru tua, baris pertama dibekukan, filter otomatis."""
    if ws.max_row < 1 or ws.max_column < 1:
        return
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def autofit_worksheet(ws: Worksheet, max_width: int = 60) -> None:
    """Lebarkan tiap kolom sesuai konten terpanjangnya (dibatasi ``max_width``)."""
    widths: Dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            length = len(str(cell.value)) + 2
            widths[cell.column] = min(max(widths.get(cell.column, 0), length), max_width)
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def autofit_columns(path: "str | Path", max_width: int = 60, style: bool = True) -> None:
    """Terapkan autofit (dan gaya header bila ``style``) pada semua sheet sebuah file."""
    wb = load_workbook(path)
    for ws in wb.worksheets:
        if style:
            style_header(ws)
        autofit_worksheet(ws, max_width)
    wb.save(path)
