"""
Ekspor hasil pipeline ke satu file Excel multi-sheet.

Sheet ditulis kondisional, hanya dibuat jika datanya tersedia, supaya file
tetap ringkas untuk run yang tidak mengaktifkan semua fitur (mis. tanpa
ADMET atau tanpa validasi RMSD). Sheet ADMET diwarnai per sel menurut
``admet.admet_rules`` (hijau/kuning/merah) dan disertai sheet legenda yang
merinci ambang batas tiap kolom.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from chemflow.admet.admet_rules import (
    CATEGORIES, EXCLUDED_COLUMNS, Flag, category_of, classify_admet_value,
    get_rule, rules_by_category, threshold_text,
)
from chemflow.utils.excel_utils import BODY_FONT, FONT_NAME, autofit_worksheet, style_header

SHEET_SUMMARY = "Ringkasan Ligan"
SHEET_DOCKING = "Hasil Docking"
SHEET_LIPINSKI = "Fisikokimia (Lipinski)"
SHEET_ADMET = "ADMET"
SHEET_ADMET_LEGEND = "Legenda ADMET"
SHEET_RMSD = "Validasi RMSD"
SHEET_REPLICATES = "Statistik Replikasi"
SHEET_GROUP_SUMMARY = "Ringkasan per Grup"
SHEET_GROUP_DOCKING = "Docking per Grup"

_FILL = {
    Flag.GREEN: PatternFill("solid", fgColor="C6EFCE"),
    Flag.YELLOW: PatternFill("solid", fgColor="FFEB9C"),
    Flag.RED: PatternFill("solid", fgColor="FFC7CE"),
}
_FONT = {
    Flag.GREEN: Font(name=FONT_NAME, size=10, color="006100"),
    Flag.YELLOW: Font(name=FONT_NAME, size=10, color="9C6500"),
    Flag.RED: Font(name=FONT_NAME, size=10, color="9C0006"),
}
_CATEGORY_HEADER_FILL = {
    "Fisikokimia": "D9E1F2",
    "Absorpsi": "E2EFDA",
    "Distribusi": "FFF2CC",
    "Metabolisme": "FCE4D6",
    "Ekskresi": "EDEDED",
    "Toksisitas": "F4CCCC",
}


def export_results(
    output_path: "str | Path",
    *,
    ligand_summary_rows: List[Dict[str, Any]],
    docking_rows: List[Dict[str, Any]],
    lipinski_rows: Optional[List[Dict[str, Any]]] = None,
    admet_rows: Optional[List[Dict[str, Any]]] = None,
    rmsd_rows: Optional[List[Dict[str, Any]]] = None,
    replicate_stats_rows: Optional[List[Dict[str, Any]]] = None,
    group_summary_rows: Optional[List[Dict[str, Any]]] = None,
    group_docking_rows: Optional[List[Dict[str, Any]]] = None,
    logger: Optional[logging.Logger] = None,
) -> Path:
    """Tulis seluruh hasil pipeline ke satu workbook Excel multi-sheet.

    Args:
        output_path: path file .xlsx keluaran.
        ligand_summary_rows: ringkasan per ligan (nama, SMILES, grup, status).
        docking_rows: satu baris per pasangan (ligan, reseptor, replikat).
        lipinski_rows: sifat fisikokimia Lipinski Ro5 per ligan.
        admet_rows: hasil ADMET per ligan (kolom persis seperti CSV ADMETLab3).
        rmsd_rows: hasil validasi RMSD redocking per reseptor.
        replicate_stats_rows: agregasi mean/std afinitas antar replikat.
        group_summary_rows: ringkasan per grup (fisikokimia, % lolos Ro5, skor kategori ADMET).
        group_docking_rows: ΔG per reseptor x grup, dibanding ligan native.
        logger: logger opsional.

    Returns:
        Path file Excel yang ditulis.
    """
    log = logger or logging.getLogger(__name__)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(str(output_path), engine="openpyxl") as writer:
        _write_sheet(writer, SHEET_SUMMARY, ligand_summary_rows)
        _write_sheet(writer, SHEET_DOCKING, docking_rows)
        if lipinski_rows:
            _write_sheet(writer, SHEET_LIPINSKI, lipinski_rows)
        if admet_rows:
            _write_sheet(writer, SHEET_ADMET, admet_rows, column_order=_admet_column_order(admet_rows))
            _write_sheet(writer, SHEET_ADMET_LEGEND, _admet_legend_rows())
        if rmsd_rows:
            _write_sheet(writer, SHEET_RMSD, rmsd_rows)
        if replicate_stats_rows:
            _write_sheet(writer, SHEET_REPLICATES, replicate_stats_rows)
        if group_summary_rows:
            _write_sheet(writer, SHEET_GROUP_SUMMARY, group_summary_rows)
        if group_docking_rows:
            _write_sheet(writer, SHEET_GROUP_DOCKING, group_docking_rows)

    _finalize_workbook(output_path)
    log.info(f"Hasil diekspor ke: {output_path}")
    return output_path


def _write_sheet(writer: pd.ExcelWriter, sheet_name: str, rows: List[Dict[str, Any]],
                 column_order: Optional[List[str]] = None) -> None:
    df = pd.DataFrame(rows) if rows else pd.DataFrame()
    if column_order:
        df = df[[c for c in column_order if c in df.columns]]
    df.to_excel(writer, sheet_name=sheet_name[:31], index=False)


def _admet_column_order(rows: List[Dict[str, Any]]) -> List[str]:
    """Urutan kolom sheet ADMET: identitas, lalu per kategori, kolom tak dikenal di akhir."""
    present = list(rows[0].keys())
    for row in rows[1:]:
        present.extend(k for k in row.keys() if k not in present)

    identity = [c for c in ("ligand", "group", "raw_smiles", "smiles") if c in present]
    ordered = list(identity)
    for category in CATEGORIES:
        ordered.extend(r.column for r in rules_by_category(category) if r.column in present and r.column not in ordered)
    for column in present:
        rule = get_rule(column)
        if column not in ordered and (rule is None or rule.column not in ordered):
            ordered.append(column)
    return ordered


def _admet_legend_rows() -> List[Dict[str, Any]]:
    rows = []
    for category in CATEGORIES:
        for rule in rules_by_category(category):
            rows.append({
                "Kategori": category,
                "Kolom": rule.column,
                "Label": rule.label,
                "Satuan": rule.unit,
                "Ambang": threshold_text(rule),
                "Rujukan": rule.citation,
                "Catatan": rule.note,
            })
    return rows


def _finalize_workbook(path: Path) -> None:
    wb = load_workbook(path)
    for ws in wb.worksheets:
        style_header(ws)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = BODY_FONT
        if ws.title == SHEET_ADMET:
            _color_admet_sheet(ws)
        elif ws.title == SHEET_RMSD:
            _color_rmsd_sheet(ws)
        autofit_worksheet(ws)
    wb.save(path)


def _color_admet_sheet(ws) -> None:
    headers = [c.value for c in ws[1]]
    for col_idx, name in enumerate(headers, start=1):
        if name is None or name in EXCLUDED_COLUMNS:
            continue
        category = category_of(str(name))
        header_fill = _CATEGORY_HEADER_FILL.get(category)
        if header_fill:
            header_cell = ws.cell(row=1, column=col_idx)
            header_cell.fill = PatternFill("solid", fgColor=header_fill)
            header_cell.font = Font(name=FONT_NAME, size=10, bold=True, color="1F1F1F")
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            flag = classify_admet_value(str(name), cell.value).flag
            if flag in _FILL:
                cell.fill = _FILL[flag]
                cell.font = _FONT[flag]


def _color_rmsd_sheet(ws) -> None:
    headers = [c.value for c in ws[1]]
    if "status" not in headers:
        return
    status_idx = headers.index("status") + 1
    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=status_idx)
        text = str(cell.value or "").lower()
        if "good" in text or "baik" in text:
            flag = Flag.GREEN
        elif "acceptable" in text or "cukup" in text:
            flag = Flag.YELLOW
        elif "poor" in text or "buruk" in text:
            flag = Flag.RED
        else:
            continue
        cell.fill = _FILL[flag]
        cell.font = _FONT[flag]
