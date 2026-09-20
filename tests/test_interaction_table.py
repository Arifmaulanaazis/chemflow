"""Test tabel Non-bond BIOVIA: TSV hasil salin menjadi Excel yang terbaca pembaca similaritas."""

from openpyxl import load_workbook

from chemflow.interaction.table import DEFAULT_HEADERS, NonbondTable, parse_rows, write_interaction_xlsx
from chemflow.similarity.biovia_reader import filter_protein_ligand, read_biovia_interactions

# Baris asli Discovery Studio 2021: 17 kolom tanpa judul, pi tanpa nama atom, kolom sudut sebagian kosong.
TSV = (
    "A:TYR334:OH - X:ASP1:O8\tYes\t0 255 0\tLigand Non-bond Monitor\t2.97191\tHydrogen Bond\t"
    "Conventional Hydrogen Bond\tA:TYR334:OH\tH-Donor\tX:ASP1:O8\tH-Acceptor\t110.59\t16.341\t\t\t\t\r\n"
    "X:ASP1:C1 - A:PHE330\tYes\t200 100 255\tLigand Non-bond Monitor\t3.66501\tHydrophobic\t"
    "Pi-Sigma\tX:ASP1:C1\tC-H\tA:PHE330\tPi-Orbitals\t\t\t\t\t14.885\t13.518\r\n"
)


def test_parse_rows_crlf_dan_baris_kosong():
    rows = parse_rows(TSV + "\r\n\r\n")
    assert len(rows) == 2
    assert all(len(row) == 17 for row in rows)
    assert rows[1][9] == "A:PHE330"


def test_parse_rows_lf_dan_kosong():
    assert parse_rows("a\tb\nc\td\n") == [["a", "b"], ["c", "d"]]
    assert parse_rows("") == []
    assert parse_rows("\t\t\n") == []


def test_write_xlsx_judul_bawaan_dan_angka(tmp_path):
    path = tmp_path / "x.xlsx"
    assert write_interaction_xlsx(NonbondTable(DEFAULT_HEADERS, TSV), path) == 2

    sheet = load_workbook(path).active
    assert sheet.title == "Non-bond"
    assert [c.value for c in sheet[1]] == list(DEFAULT_HEADERS)
    first = [c.value for c in sheet[2]]
    assert first[4] == 2.97191 and first[11] == 110.59       # jarak dan sudut jadi angka
    assert first[5] == "Hydrogen Bond" and first[7] == "A:TYR334:OH"
    assert first[13] is None                                   # sel kosong tetap kosong
    assert sheet.freeze_panes == "A2"


def test_write_xlsx_judul_dari_antarmuka_dan_kolom_lebih_banyak(tmp_path):
    path = tmp_path / "x.xlsx"
    write_interaction_xlsx(NonbondTable(("Name", "Visible"), "a\tYes\tx\ty\n"), path)
    header = [c.value for c in load_workbook(path).active[1]]
    assert header == ["Name", "Visible", "Kolom 3", "Kolom 4"]


def test_write_xlsx_tanpa_interaksi_hanya_judul(tmp_path):
    path = tmp_path / "x.xlsx"
    assert write_interaction_xlsx(NonbondTable(DEFAULT_HEADERS, ""), path) == 0
    sheet = load_workbook(path).active
    assert sheet.max_row == 1
    assert read_biovia_interactions(path) == []               # pembaca menerima berkas berisi judul saja


def test_write_xlsx_nilai_berawalan_sama_dengan_tetap_teks(tmp_path):
    path = tmp_path / "x.xlsx"
    write_interaction_xlsx(NonbondTable(("Name",), "=1+1\n"), path)
    cell = load_workbook(path).active["A2"]
    assert cell.value == "=1+1" and cell.data_type == "s"


def test_hasil_ekspor_terbaca_pembaca_similaritas_termasuk_interaksi_pi(tmp_path):
    path = tmp_path / "x.xlsx"
    write_interaction_xlsx(NonbondTable(DEFAULT_HEADERS, TSV), path)
    interactions = read_biovia_interactions(path)
    assert [i.type_subtype for i in interactions] == ["Conventional Hydrogen Bond", "Pi-Sigma"]
    contacts = filter_protein_ligand(interactions, ligand_chain="X")
    assert [(c.resnum, c.resname) for c in contacts] == [(334, "TYR"), (330, "PHE")]
