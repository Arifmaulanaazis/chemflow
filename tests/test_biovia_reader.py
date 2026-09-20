"""Test biovia_reader: parser export interaksi BIOVIA Discovery Studio."""

from pathlib import Path

import pytest
from openpyxl import Workbook

from chemflow.similarity.biovia_reader import (
    filter_protein_ligand, parse_atom_spec, read_biovia_interactions,
)


def _write_biovia_excel(path: Path, rows, with_header: bool = False) -> None:
    wb = Workbook()
    ws = wb.active
    if with_header:
        ws.append(["Name", "Rendered", "Color", "Style", "ID", "Category", "Type",
                   "From", "From Chemistry", "To", "To Chemistry"])
    for row in rows:
        ws.append(row)
    wb.save(path)


def _row(from_spec, from_chem, to_spec, to_chem, category="Hydrogen Bond",
         type_subtype="Conventional Hydrogen Bond"):
    return [f"{from_spec} - {to_spec}", "Yes", "0 255 0", "Ligand Non-bond Monitor", 1,
            category, type_subtype, from_spec, from_chem, to_spec, to_chem]


def test_parse_atom_spec_valid():
    spec = parse_atom_spec("A:SER195:OG")
    assert spec.chain == "A"
    assert spec.resname == "SER"
    assert spec.resnum == 195
    assert spec.atom_name == "OG"


def test_parse_atom_spec_ligand_chain():
    spec = parse_atom_spec("X:ASP1:O8")
    assert spec.chain == "X"
    assert spec.resname == "ASP"
    assert spec.resnum == 1


def test_parse_atom_spec_resname_mengandung_angka():
    spec = parse_atom_spec("X:C6X1:O2")
    assert (spec.chain, spec.resname, spec.resnum, spec.atom_name) == ("X", "C6X", 1, "O2")
    spec = parse_atom_spec("C:02J1:C1")
    assert (spec.resname, spec.resnum) == ("02J", 1)


def test_parse_atom_spec_resnum_negatif_dan_ion():
    assert parse_atom_spec("A:GLY-5:CA").resnum == -5
    spec = parse_atom_spec("A:CL1000:CL")
    assert (spec.resname, spec.resnum) == ("CL", 1000)


def test_parse_atom_spec_invalid_raise():
    with pytest.raises(ValueError):
        parse_atom_spec("bukan spek atom")


def test_read_biovia_interactions_tanpa_header(tmp_path):
    path = tmp_path / "interaksi.xlsx"
    rows = [
        _row("A:SER195:OG", "H-Donor", "X:ASP1:O8", "H-Acceptor"),
        _row("A:GLY216:N", "H-Donor", "X:ASP1:O14", "H-Acceptor"),
    ]
    _write_biovia_excel(path, rows, with_header=False)

    interactions = read_biovia_interactions(path)
    assert len(interactions) == 2
    assert interactions[0].from_spec.resname == "SER"
    assert interactions[0].to_spec.chain == "X"


def test_read_biovia_interactions_dengan_header_dilewati(tmp_path):
    path = tmp_path / "interaksi.xlsx"
    rows = [_row("A:SER195:OG", "H-Donor", "X:ASP1:O8", "H-Acceptor")]
    _write_biovia_excel(path, rows, with_header=True)

    interactions = read_biovia_interactions(path)
    assert len(interactions) == 1  # baris header tidak dihitung sbg data


def test_read_biovia_interactions_kolom_trailing_hilang_tetap_jalan(tmp_path):
    """Kolom setelah 'To Chemistry' boleh tidak ada, file tetap terbaca."""
    path = tmp_path / "interaksi.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["A:SER195:OG - X:ASP1:O8", "Yes", "0 255 0", "Ligand Non-bond Monitor", 1,
               "Hydrogen Bond", "Conventional Hydrogen Bond", "A:SER195:OG", "H-Donor",
               "X:ASP1:O8", "H-Acceptor"])  # persis 11 kolom, tanpa Distance/Angle
    wb.save(path)

    interactions = read_biovia_interactions(path)
    assert len(interactions) == 1


def test_read_biovia_interactions_kolom_kurang_dari_11_raise(tmp_path):
    path = tmp_path / "interaksi.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Yes", "0 255 0", "Ligand Non-bond Monitor", 1,
               "Hydrogen Bond", "Conventional Hydrogen Bond", "A:SER195:OG"])  # cuma 8 kolom
    wb.save(path)

    with pytest.raises(ValueError, match="minimal 11"):
        read_biovia_interactions(path)


def test_read_biovia_interactions_file_tidak_ada_raise(tmp_path):
    with pytest.raises(FileNotFoundError):
        read_biovia_interactions(tmp_path / "tidak_ada.xlsx")


def test_filter_protein_ligand_buang_intra_ligand_dan_intra_protein(tmp_path):
    path = tmp_path / "interaksi.xlsx"
    rows = [
        _row("A:SER195:OG", "H-Donor", "X:ASP1:O8", "H-Acceptor"),          # protein-ligan, disimpan
        _row("X:ASP1:H9", "H-Donor", "X:ASP1:O14", "H-Acceptor"),          # intra-ligand, dibuang
        _row("A:CYS191:O", "H-Acceptor", "A:SER195:OG", "H-Donor"),        # intra-protein, dibuang
    ]
    _write_biovia_excel(path, rows)

    interactions = read_biovia_interactions(path)
    contacts = filter_protein_ligand(interactions, ligand_chain="X")

    assert len(contacts) == 1
    assert contacts[0].resname == "SER"
    assert contacts[0].resnum == 195


def test_filter_protein_ligand_sisi_ligan_bisa_dari_atau_ke():
    from chemflow.similarity.biovia_reader import AtomSpec, BiovaInteraction

    it_from_ligand = BiovaInteraction(
        category="Hydrophobic", type_subtype="Pi-Alkyl",
        from_spec=AtomSpec("X", "ASP", 1, "C1"), from_chemistry="Pi-Alkyl",
        to_spec=AtomSpec("A", "ALA", 50, "CB"), to_chemistry="Alkyl",
    )
    contacts = filter_protein_ligand([it_from_ligand], ligand_chain="X")
    assert len(contacts) == 1
    assert contacts[0].resname == "ALA"
    assert contacts[0].resnum == 50


def test_parse_atom_spec_interaksi_pi_tanpa_nama_atom():
    spec = parse_atom_spec("A:PHE330")
    assert (spec.chain, spec.resname, spec.resnum, spec.atom_name) == ("A", "PHE", 330, "")


def test_read_biovia_interactions_kolom_dicari_lewat_judul(tmp_path):
    """Urutan dan kolom tambahan bebas selama judul From, To, Category, Types ada."""
    path = tmp_path / "interaksi.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Distance", "To", "From", "Types", "Category", "To Chemistry", "From Chemistry", "Name"])
    ws.append([2.9, "X:LIG1:O8", "A:SER195:OG", "Conventional Hydrogen Bond", "Hydrogen Bond",
               "H-Acceptor", "H-Donor", "x"])
    wb.save(path)

    (interaction,) = read_biovia_interactions(path)
    assert interaction.from_spec.resname == "SER" and interaction.to_spec.chain == "X"
    assert (interaction.category, interaction.type_subtype) == ("Hydrogen Bond", "Conventional Hydrogen Bond")
    assert (interaction.from_chemistry, interaction.to_chemistry) == ("H-Donor", "H-Acceptor")


def test_read_biovia_interactions_hanya_judul_berarti_tanpa_interaksi(tmp_path):
    path = tmp_path / "interaksi.xlsx"
    _write_biovia_excel(path, [], with_header=True)
    assert read_biovia_interactions(path) == []


def test_read_biovia_interactions_satu_baris_bukan_judul_dan_bukan_data_tetap_raise(tmp_path):
    path = tmp_path / "interaksi.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["a", "b", "c", "d", "e", "f", "g", "h", "i", "j", "k"])
    wb.save(path)
    with pytest.raises(ValueError, match="Format spek atom"):
        read_biovia_interactions(path)


def test_interaksi_pi_ikut_terhitung_sebagai_kontak_protein(tmp_path):
    path = tmp_path / "interaksi.xlsx"
    rows = [_row("X:LIG1:C1", "C-H", "A:PHE330", "Pi-Orbitals", category="Hydrophobic", type_subtype="Pi-Sigma")]
    _write_biovia_excel(path, rows)
    (contact,) = filter_protein_ligand(read_biovia_interactions(path), ligand_chain="X")
    assert (contact.resnum, contact.resname, contact.interaction_type) == (330, "PHE", "Pi-Sigma")
