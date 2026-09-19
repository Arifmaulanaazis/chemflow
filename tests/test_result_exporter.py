"""Test export_results: sheet kondisional, pewarnaan ADMET, legenda, gaya header."""

import openpyxl

from chemflow.io.result_exporter import export_results


def _basic(out, **extra):
    return export_results(
        out,
        ligand_summary_rows=[{"name": "A", "status": "siap"}],
        docking_rows=[{"ligand": "A", "receptor": "R1", "affinity_best": -6.0}],
        **extra,
    )


def _fill(cell):
    return cell.fill.start_color.rgb[-6:]


def test_output_path_string_diterima(tmp_path):
    out = str(tmp_path / "hasil.xlsx")
    _basic(out)
    assert "Ringkasan Ligan" in openpyxl.load_workbook(out).sheetnames


def test_sheet_wajib_selalu_ada_dan_opsional_tidak(tmp_path):
    out = tmp_path / "hasil.xlsx"
    _basic(out)
    wb = openpyxl.load_workbook(out)
    assert "Ringkasan Ligan" in wb.sheetnames
    assert "Hasil Docking" in wb.sheetnames
    assert "ADMET" not in wb.sheetnames
    assert "Legenda ADMET" not in wb.sheetnames


def test_sheet_opsional_muncul_kalau_data_ada(tmp_path):
    out = tmp_path / "hasil.xlsx"
    _basic(
        out,
        lipinski_rows=[{"ligand": "A", "MW": 180.0}],
        admet_rows=[{"ligand": "A", "Lipinski": 0}],
        rmsd_rows=[{"receptor": "R1", "rmsd_angstrom": 0.5, "status": "good"}],
        replicate_stats_rows=[{"ligand": "A", "affinity_mean": -6.0}],
    )
    wb = openpyxl.load_workbook(out)
    for expected in ("Fisikokimia (Lipinski)", "ADMET", "Legenda ADMET", "Validasi RMSD", "Statistik Replikasi"):
        assert expected in wb.sheetnames


def test_header_bergaya_beku_dan_berfilter(tmp_path):
    out = tmp_path / "hasil.xlsx"
    _basic(out)
    ws = openpyxl.load_workbook(out)["Hasil Docking"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None
    assert ws["A1"].font.bold is True


def test_pewarnaan_admet_hijau_kuning_merah(tmp_path):
    out = tmp_path / "hasil.xlsx"
    _basic(out, admet_rows=[
        {"ligand": "A", "hERG": 0.1, "DILI": 0.5, "Ames": 0.9, "PAINS": "['-']", "caco2": -4.0},
    ])
    ws = openpyxl.load_workbook(out)["ADMET"]
    headers = {c.value: c.column for c in ws[1]}
    assert _fill(ws.cell(2, headers["hERG"])) == "C6EFCE"
    assert _fill(ws.cell(2, headers["DILI"])) == "FFEB9C"
    assert _fill(ws.cell(2, headers["Ames"])) == "FFC7CE"
    assert _fill(ws.cell(2, headers["PAINS"])) == "C6EFCE"
    assert _fill(ws.cell(2, headers["caco2"])) == "C6EFCE"


def test_pewarnaan_admet_kolom_tak_dikenal_dan_identitas_tak_diwarnai(tmp_path):
    out = tmp_path / "hasil.xlsx"
    _basic(out, admet_rows=[{"ligand": "A", "smiles": "CCO", "KolomAneh": 42, "hERG": 0.1}])
    ws = openpyxl.load_workbook(out)["ADMET"]
    headers = {c.value: c.column for c in ws[1]}
    assert _fill(ws.cell(2, headers["KolomAneh"])) != "C6EFCE"
    assert _fill(ws.cell(2, headers["smiles"])) != "C6EFCE"


def test_kolom_admet_diurut_per_kategori_kolom_asing_di_akhir(tmp_path):
    out = tmp_path / "hasil.xlsx"
    _basic(out, admet_rows=[{"ligand": "A", "KolomAneh": 1, "hERG": 0.1, "hia": 0.2, "MW": 300}])
    ws = openpyxl.load_workbook(out)["ADMET"]
    headers = [c.value for c in ws[1]]
    assert headers[0] == "ligand"
    assert headers.index("MW") < headers.index("hia") < headers.index("hERG") < headers.index("KolomAneh")


def test_legenda_admet_memuat_ambang(tmp_path):
    out = tmp_path / "hasil.xlsx"
    _basic(out, admet_rows=[{"ligand": "A", "hERG": 0.1}])
    ws = openpyxl.load_workbook(out)["Legenda ADMET"]
    headers = [c.value for c in ws[1]]
    assert headers == ["Kategori", "Kolom", "Label", "Satuan", "Ambang", "Rujukan", "Catatan"]
    columns = [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]
    assert "hERG" in columns and "caco2" in columns


def test_pewarnaan_rmsd_status(tmp_path):
    out = tmp_path / "hasil.xlsx"
    _basic(out, rmsd_rows=[{"receptor": "R1", "status": "good"}, {"receptor": "R2", "status": "poor"}])
    ws = openpyxl.load_workbook(out)["Validasi RMSD"]
    fills = [_fill(cell) for row in ws.iter_rows(min_row=2) for cell in row
             if str(cell.value).lower() in ("good", "poor")]
    assert fills == ["C6EFCE", "FFC7CE"]
