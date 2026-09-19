"""Test generator template Excel: file yang ditulis harus benar-benar bisa
dibaca oleh io.excel_ligands/io.excel_receptors chemflow sendiri, bukan cuma
"terlihat benar"."""

from chemflow.io.excel_ligands import read_ligands
from chemflow.io.excel_receptors import read_receptors
from chemflow.templates import (
    write_all_templates, write_ligand_tidy_template, write_ligand_wide_template, write_receptor_template,
)


def test_template_tidy_terbaca_sebagai_mode_tidy(tmp_path):
    path = write_ligand_tidy_template(tmp_path / "ligan.xlsx")
    assert path.exists()

    records = read_ligands(path)
    assert len(records) == 4
    names = {r.name for r in records}
    assert "Quercetin" in names
    aspirin = next(r for r in records if r.name == "Aspirin")
    assert aspirin.needs_pubchem_lookup is True  # SMILES sengaja dikosongkan


def test_template_wide_terbaca_sebagai_mode_wide(tmp_path):
    path = write_ligand_wide_template(tmp_path / "ligan.xlsx")
    records = read_ligands(path)
    groups = {r.group for r in records}
    assert groups == {"Tanaman_X", "Tanaman_Y", "Tanaman_Z"}
    assert all(r.needs_pubchem_lookup for r in records)


def test_template_reseptor_terbaca_dengan_multi_situs_dan_multi_pdb(tmp_path):
    path = write_receptor_template(tmp_path / "reseptor.xlsx")
    entries = read_receptors(path)

    codes = [e.pdb_code for e in entries]
    assert codes.count("3PTB") == 2  # multi-situs
    assert "6LU7" in codes and "1AKI" in codes  # multi-reseptor

    ptb_entries = [e for e in entries if e.pdb_code == "3PTB"]
    assert all(e.has_gridbox_center for e in ptb_entries)
    assert ptb_entries[0].unique_key != ptb_entries[1].unique_key


def test_write_all_templates_menghasilkan_3_file(tmp_path):
    paths = write_all_templates(tmp_path / "out")
    assert len(paths) == 3
    assert all(p.exists() for p in paths)


def test_template_punya_sheet_petunjuk_kedua(tmp_path):
    import openpyxl

    path = write_ligand_tidy_template(tmp_path / "ligan.xlsx")
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames[0] != "Petunjuk"  # sheet data harus di posisi pertama
    assert "Petunjuk" in wb.sheetnames
